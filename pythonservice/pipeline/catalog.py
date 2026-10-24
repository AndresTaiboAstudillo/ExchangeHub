"""Catálogo oficial del Mineduc: descubrimiento de la edición, lectura, validación y limpieza."""

import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from urllib.parse import urljoin

import openpyxl
from pydantic import BaseModel, ValidationError, field_validator

from .errors import SourceFormatError
from .models import Provenance, Rejected, UniversityRecord
from .text import collapse_spaces, display_name, name_key, strip_accents

SOURCE = "mineduc_ies"
_FILE_RE = re.compile(r"""href=["']([^"']*ies_vig_([a-z]{3})_(\d{2})\.xlsx)["']""", re.IGNORECASE)
_MONTHS = {
    m: i + 1
    for i, m in enumerate(
        ["ene", "feb", "mar", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic"]
    )
}
_EDITION_RE = re.compile(r"ies_vig_([a-z]{3})_(\d{2})\.xlsx", re.IGNORECASE)
HEADER_SCAN_ROWS = 10

# Cabecera normalizada (sin tildes, minúsculas) → campo interno
_HEADERS = {
    "numero registro": "registry_number",
    "nombre": "name",
    "tipo institucion": "institution_type",
    "vigencia": "vigencia",
    "funcionamiento": "operating_status",
}

_TYPES: dict[str, tuple[str, str]] = {
    "universidad estatal": ("universidad", "estatal"),
    "universidad particular con aportes": ("universidad", "particular_con_aportes"),
    "universidad privada": ("universidad", "privada"),
    "instituto profesional": ("instituto_profesional", "otra"),
    "centro de formacion tecnica": ("centro_formacion_tecnica", "otra"),
    "centro de formacion tecnica estatal": ("centro_formacion_tecnica", "estatal"),
}


def _edition(mon: str, yy: str) -> tuple[str, int] | None:
    month = _MONTHS.get(strip_accents(mon).lower())
    if month is None:
        return None
    year = 2000 + int(yy)
    return f"{year}-{month:02d}", year * 100 + month


def edition_from_url(url: str) -> str | None:
    m = _EDITION_RE.search(url)
    if not m:
        return None
    ed = _edition(m.group(1), m.group(2))
    return ed[0] if ed else None


def discover_latest_url(page_html: str, page_url: str) -> tuple[str, str] | None:
    """Busca en la ficha el archivo ies_vig_<mes>_<aa>.xlsx más reciente."""
    best: tuple[int, str, str] | None = None
    for m in _FILE_RE.finditer(page_html):
        ed = _edition(m.group(2), m.group(3))
        if ed is None:
            continue
        candidate = (ed[1], urljoin(page_url, m.group(1)), ed[0])
        if best is None or candidate[0] > best[0]:
            best = candidate
    return (best[1], best[2]) if best else None


class CatalogRow(BaseModel):
    registry_number: str
    name: str
    institution_type: str
    vigencia: str
    operating_status: str

    @field_validator("*", mode="before")
    @classmethod
    def _to_text(cls, v):
        if v is None:
            raise ValueError("campo vacío")
        text = collapse_spaces(str(v))
        if not text:
            raise ValueError("campo vacío")
        return text

    @field_validator("registry_number")
    @classmethod
    def _registry(cls, v: str) -> str:
        v = v.upper()
        if not re.fullmatch(r"[A-Z]\d{1,4}", v):
            raise ValueError(f"número de registro con formato inválido: {v!r}")
        return v


@dataclass
class CatalogParse:
    records: list[UniversityRecord] = field(default_factory=list)
    rejected: list[Rejected] = field(default_factory=list)
    rows_received: int = 0
    duplicates_discarded: int = 0


def _find_header(rows: list[tuple]) -> tuple[int, dict[int, str]]:
    for idx, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        for col, cell in enumerate(row):
            if cell is None:
                continue
            key = collapse_spaces(strip_accents(str(cell)).lower())
            if key in _HEADERS:
                mapping[col] = _HEADERS[key]
        if set(mapping.values()) == set(_HEADERS.values()):
            return idx, mapping
    raise SourceFormatError(
        SOURCE, "no se encontró la fila de cabecera esperada (NÚMERO REGISTRO, NOMBRE, ...)"
    )


def parse_catalog(content: bytes, *, fetched_at: datetime, url: str) -> CatalogParse:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    except Exception as exc:  # archivo corrupto o no es xlsx
        raise SourceFormatError(SOURCE, f"no se pudo leer el Excel ({type(exc).__name__})") from exc

    header_idx, mapping = _find_header(rows)
    result = CatalogParse()
    seen_numbers: set[str] = set()
    seen_names: set[str] = set()

    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue  # fila en blanco: no cuenta como registro
        result.rows_received += 1
        raw = {mapping[c]: row[c] if c < len(row) else None for c in mapping}
        try:
            parsed = CatalogRow(**raw)
        except ValidationError as exc:
            msgs = "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors())
            ext = str(raw.get("registry_number") or "") or None
            result.rejected.append(
                Rejected(source=SOURCE, row=offset, external_id=ext, reason=msgs)
            )
            continue

        type_key = name_key(parsed.institution_type)
        classified = _TYPES.get(type_key)
        if classified is None:
            result.rejected.append(
                Rejected(
                    source=SOURCE,
                    row=offset,
                    external_id=parsed.registry_number,
                    reason=f"tipo de institución desconocido: {parsed.institution_type!r}",
                )
            )
            continue

        nkey = name_key(parsed.name)
        if parsed.registry_number in seen_numbers or nkey in seen_names:
            result.duplicates_discarded += 1
            continue
        seen_numbers.add(parsed.registry_number)
        seen_names.add(nkey)

        row_hash = hashlib.sha256(
            "|".join(str(raw[k]) for k in sorted(raw)).encode("utf-8")
        ).hexdigest()
        category, ownership = classified
        result.records.append(
            UniversityRecord(
                registry_number=parsed.registry_number,
                name=display_name(parsed.name),
                name_key=nkey,
                institution_type=parsed.institution_type,
                category=category,
                ownership=ownership,
                operating_status=name_key(parsed.operating_status).replace(" ", "_"),
                is_vigente=name_key(parsed.vigencia) == "vigente",
                provenance=[
                    Provenance(
                        source=SOURCE,
                        external_id=parsed.registry_number,
                        url=url,
                        fetched_at=fetched_at,
                        content_hash=row_hash,
                    )
                ],
            )
        )
    return result
